"""
Dumps an arbitrary (non-template) Excel workbook into a compact, bounded
text representation for the Groq-based extractor (see
app/services/extraction_service.py). Every sheet, every non-empty row and
cell - capped so a large or malicious workbook can't blow up the LLM's
context window or this platform's outbound request size. Reuses the same
openpyxl reader as the strict template parser (excel_parser.py); this is
just a different, layout-agnostic read of the same file.
"""
import io

import openpyxl

from app.core.exceptions import IntakeParseError

_MAX_SHEETS = 10
_MAX_ROWS_PER_SHEET = 300
_MAX_CHARS = 24_000


def dump_workbook_text(file_bytes: bytes) -> str:
    try:
        wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True, read_only=True)
    except Exception as exc:  # openpyxl raises several distinct types for corrupt files
        raise IntakeParseError(
            "The uploaded file is not a valid Excel (.xlsx) workbook, or could not be opened."
        ) from exc

    lines: list[str] = []
    for sheet_name in wb.sheetnames[:_MAX_SHEETS]:
        ws = wb[sheet_name]
        sheet_lines: list[str] = []
        for row in ws.iter_rows():
            cells = [str(c.value).strip() for c in row if c.value is not None and str(c.value).strip() != ""]
            if not cells:
                continue
            sheet_lines.append(" | ".join(cells))
            if len(sheet_lines) >= _MAX_ROWS_PER_SHEET:
                sheet_lines.append("... (truncated)")
                break
        if not sheet_lines:
            continue
        lines.append(f"### Sheet: {sheet_name}")
        lines.extend(sheet_lines)

    return "\n".join(lines)[:_MAX_CHARS]
